# -*- coding: utf-8 -*-
"""Mirror every table of the pipeline to Excel.

Each CSV listed below gets an .xlsx next to it with the same name: one sheet,
frozen and filterable header, readable column widths. The CSV stays the source
of truth (every script reads and verifies the CSVs); the workbooks are for
reading. Run last, after every table has been rebuilt:

    python3 export_tables_xlsx.py
"""
import csv
import glob
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

SC = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SC)
COUNTS = os.path.join(ROOT, "per_tool_counts_from_scratch_v2", "counts_per_tool_v2")

PATTERNS = [
    os.path.join(ROOT, "table_S*.csv"),            # supplementary tables
    os.path.join(SC, "*.csv"),                     # pipeline and curation tables
    os.path.join(SC, "r_general_scheme", "*.csv"),  # what the R figure is handed
    os.path.join(COUNTS, "*.csv"),                 # per-tool counts
]
EXCEL_CELL_LIMIT = 32767


def tables() -> list:
    found = []
    for pattern in PATTERNS:
        for path in sorted(glob.glob(pattern)):
            if not os.path.basename(path).startswith("_"):  # scratch files
                found.append(path)
    return found


def _cell(value: str):
    """Numbers become numbers in Excel; everything else stays text."""
    try:
        number = float(value)
    except ValueError:
        return value
    if number != number or number in (float("inf"), float("-inf")):  # nan, inf
        return value
    return int(number) if number.is_integer() and "." not in value and "e" not in value.lower() else number


def write_xlsx(csv_path: str) -> int:
    """Copy the CSV grid cell for cell. The per-tool counts files open with
    '# key,value' metadata lines and a blank line; they are kept as they are,
    and the header and filter go on the first row of the table below them."""
    with open(csv_path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    header_row = 1
    while header_row <= len(rows) and (not rows[header_row - 1]
                                       or rows[header_row - 1][0].startswith("#")):
        header_row += 1
    wb = Workbook()
    ws = wb.active
    ws.title = os.path.splitext(os.path.basename(csv_path))[0][:31]
    widths = {}
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if len(value) > EXCEL_CELL_LIMIT:
                sys.exit(f"{csv_path}: a cell is longer than Excel allows (row {r}, column {c})")
            ws.cell(r, c, _cell(value) if r > header_row else value)
            if r >= header_row and r < header_row + 500:
                widths[c] = max(widths.get(c, 0), len(value))
    if header_row <= len(rows):
        ws.freeze_panes = ws.cell(header_row + 1, 1)
        last = ws.cell(len(rows), max(len(rows[header_row - 1]), 1))
        ws.auto_filter.ref = f"A{header_row}:{last.coordinate}"
        for c, bold in enumerate(rows[header_row - 1], start=1):
            ws.cell(header_row, c).font = Font(bold=True)
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 8), 60)
    wb.save(os.path.splitext(csv_path)[0] + ".xlsx")
    return len(rows)


def main() -> None:
    paths = tables()
    for path in paths:
        write_xlsx(path)
    print(f"{len(paths)} tables mirrored to .xlsx")


if __name__ == "__main__":
    main()
